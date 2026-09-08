import hashlib
import io
import json
from pathlib import Path
from types import SimpleNamespace
from zipfile import ZIP_DEFLATED, ZipFile

import pytest

CORPUS = Path(__file__).resolve().parents[2] / "docs/evaluation/knowledge"
MIMES = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "csv": "text/csv",
    "md": "text/markdown",
    "txt": "text/plain",
}


def parse(data, kind="txt", **overrides):
    from shopsteward_knowledge.parsing import parse_original

    fields = dict(
        version_id="V1",
        sha256=hashlib.sha256(data).hexdigest(),
        size_bytes=len(data),
        mime=MIMES.get(kind, kind),
        storage_key=f"originals/file.{kind}",
    )
    return parse_original(SimpleNamespace(**(fields | overrides)), data)


def test_utf8_text_preserves_paragraphs_and_markdown_heading_path():
    result = parse("# 说明\n\n首段\n续行\n\n## 条件\n\n末段".encode(), "md")
    assert result.status == "COMPLETE"
    assert [b["text"] for b in result.blocks] == ["# 说明", "首段\n续行", "## 条件", "末段"]
    assert result.blocks[-1]["locator"]["section_path"] == ["说明", "条件"]
    assert result.blocks[-1]["locator"]["paragraph_range"] == [4, 4]
    assert all(b["locator"]["page"] is None for b in result.blocks)
    assert result.coverage["units_unparsed"] == 0


@pytest.mark.parametrize("override", [{"sha256": "0" * 64}, {"size_bytes": 10}])
def test_original_integrity_checked_before_parsing(override):
    with pytest.raises(ValueError):
        parse(b"body", **override)


def test_rejects_invalid_utf8_without_discarding_characters():
    with pytest.raises(ValueError, match="UTF-8"):
        parse(b"before\xffafter")


def test_unknown_mime_not_promoted_by_filename():
    result = parse(b"arbitrary", "application/octet-stream", storage_key="original.pdf")
    assert result.status == "UNSUPPORTED"
    assert result.blocks == [] and result.warnings


def test_csv_preserves_header_multiline_cells_and_logical_row_locators():
    result = parse(b'name,condition\nMilk,"cold\nchain"\nEgg,fragile\n', "csv")
    assert result.status == "COMPLETE"
    assert result.blocks[-1]["locator"]["cell_range"] == "A3:B3"
    assert result.blocks[0]["table"]["headers"] == ["name", "condition"]
    assert result.blocks[0]["table"]["rows"] == [["Milk", "cold\nchain"]]
    assert "name" in result.blocks[-1]["text"] and "fragile" in result.blocks[-1]["text"]


def test_docx_tables_and_paragraphs_remain_in_actual_document_order():
    from docx import Document

    document = Document()
    document.add_heading("Conditions", level=1)
    document.add_paragraph("Before")
    table = document.add_table(rows=2, cols=2)
    for cell, value in zip(
        [c for row in table.rows for c in row.cells], ["Item", "Rule", "Milk", "Cold"], strict=True
    ):
        cell.text = value
    document.add_paragraph("After")
    stream = io.BytesIO()
    document.save(stream)
    result = parse(stream.getvalue(), "docx")
    assert result.status == "COMPLETE"
    assert [b["locator"]["kind"] for b in result.blocks] == [
        "paragraph",
        "paragraph",
        "table",
        "paragraph",
    ]
    assert result.blocks[-1]["text"] == "After"
    assert result.blocks[-1]["locator"]["paragraph_range"] == [7, 7]
    assert result.blocks[2]["table"]["headers"] == ["Item", "Rule"]
    assert result.blocks[2]["table"]["rows"] == [["Milk", "Cold"]]
    assert all(b["locator"]["page"] is None for b in result.blocks)


def workbook_bytes():
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference

    book = Workbook()
    sheet = book.active
    sheet.title = "Prices"
    sheet.append(["Item", "Total"])
    sheet.append(["Milk", "=2+3"])
    sheet.append(["Egg", "=3+4"])
    chart = BarChart()
    chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=3), titles_from_data=True)
    chart_sheet = book.create_chartsheet("Chart")
    chart_sheet.sheet_state = "hidden"
    chart_sheet.add_chart(chart)
    stream = io.BytesIO()
    book.save(stream)
    output = io.BytesIO()
    with ZipFile(stream) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(b"<f>2+3</f><v></v>", b"<f>2+3</f><v>5</v>")
            target.writestr(item, data)
    return output.getvalue()


def test_xlsx_keeps_formulas_caches_ranges_and_chartsheet_state():
    result = parse(workbook_bytes(), "xlsx")
    assert result.status == "PARTIAL"
    formulas = {f["cell"]: f for b in result.blocks for f in b.get("formulas", [])}
    assert formulas["B2"] == {
        "cell": "B2",
        "formula": "=2+3",
        "cached_value": 5,
        "cache_status": "present",
    }
    assert formulas["B3"] == {
        "cell": "B3",
        "formula": "=3+4",
        "cached_value": None,
        "cache_status": "missing",
    }
    assert result.blocks[0]["locator"]["sheet"] == "Prices"
    assert result.blocks[0]["locator"]["cell_range"] == "A2:B2"
    assert "Total" in result.blocks[0]["text"] and "=2+3" in result.blocks[0]["text"]
    chart = next(i for i in result.coverage["items"] if i.get("sheet") == "Chart")
    assert chart["kind"] == "chartsheet" and chart["state"] == "hidden"
    assert chart["status"] == "UNPARSED"


def test_chunk_ids_bind_version_profiles_location_and_text_but_not_storage_path():
    from shopsteward_knowledge.parsing.chunking import chunk_blocks

    data = b"abcdefghijklmnopqrst"
    first = parse(data)
    profile = {"id": "fixed", "max_chars": 8, "overlap_chars": 2}
    chunks = chunk_blocks(first.blocks, profile)
    assert [c["text"] for c in chunks] == ["abcdefgh", "ghijklmn", "mnopqrst"]
    assert chunks == chunk_blocks(
        parse(data, storage_key="moved/original.txt").blocks, dict(reversed(list(profile.items())))
    )
    assert len({c["id"] for c in chunks}) == 3
    for c in chunks:
        assert c["content_sha256"] == hashlib.sha256(c["text"].encode()).hexdigest()
        assert c["original_sha256"] == hashlib.sha256(data).hexdigest()
    assert chunks[0]["id"] != chunk_blocks(parse(data, version_id="V2").blocks, profile)[0]["id"]
    assert chunks[0]["id"] != chunk_blocks(first.blocks, profile | {"id": "new"})[0]["id"]
    changed = json.loads(json.dumps(first.blocks))
    changed[0]["locator"]["paragraph_range"] = [2, 2]
    assert chunks[0]["id"] != chunk_blocks(changed, profile)[0]["id"]


@pytest.mark.parametrize(
    "profile", [{"max_chars": 0}, {"max_chars": 8, "overlap_chars": 8}, {"max_chars": True}]
)
def test_invalid_chunk_budget_rejected(profile):
    from shopsteward_knowledge.parsing.chunking import chunk_blocks

    with pytest.raises(ValueError):
        chunk_blocks(parse(b"hello").blocks, profile)


def test_long_table_chunks_keep_header_and_every_data_character():
    from shopsteward_knowledge.parsing.chunking import chunk_blocks

    result = parse(b"name,rule\nMilk,abcdefghijklmnopqrstuvxyz\n", "csv")
    chunks = chunk_blocks(result.blocks, {"max_chars": 20})
    assert len(chunks) > 1
    assert all(c["text"].startswith("name\trule\n") for c in chunks)
    assert all(c["text"].startswith("name\trule\nMilk\t") for c in chunks)
    assert (
        "".join(c["text"].removeprefix("name\trule\nMilk\t") for c in chunks)
        == "abcdefghijklmnopqrstuvxyz"
    )
    assert all(len(c["text"]) <= 20 for c in chunks)


def test_review_docx_content_control_keeps_binding_exception_and_coordinates():
    from docx import Document
    from docx.oxml import OxmlElement

    document = Document()
    document.add_paragraph("Before")
    exception = document.add_paragraph("Binding exception: reject broken seals")
    wrapper = OxmlElement("w:sdt")
    content = OxmlElement("w:sdtContent")
    exception._p.addprevious(wrapper)
    wrapper.append(content)
    content.append(exception._p)
    document.add_paragraph("After")
    stream = io.BytesIO()
    document.save(stream)
    result = parse(stream.getvalue(), "docx")
    assert [b["text"] for b in result.blocks] == [
        "Before",
        "Binding exception: reject broken seals",
        "After",
    ]
    assert [b["locator"]["paragraph_range"] for b in result.blocks] == [[1, 1], [2, 2], [3, 3]]


def test_review_scanned_pdf_with_text_footer_still_requires_ocr():
    from pypdf import PdfReader, PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    document = next(d for d in DOCUMENTS if d["version_id"] == "K0-P02-v1")
    original = (CORPUS / document["source_path"]).read_bytes()
    reader = PdfReader(io.BytesIO(original))
    # Same real text-footer overlay as the reportlab repro, using the installed
    # pypdf dependency so this regression requires no PDF-generation extra.
    overlay_writer = PdfWriter()
    overlay = overlay_writer.add_blank_page(width=595, height=842)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    overlay[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
    )
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 30 20 Td (Page 2) Tj ET")
    overlay[NameObject("/Contents")] = stream
    writer = PdfWriter()
    writer.append_pages_from_reader(reader)
    writer.pages[1].merge_page(overlay)
    changed = io.BytesIO()
    writer.write(changed)
    result = parse(changed.getvalue(), "pdf")
    assert result.status == "PARTIAL"
    assert any(b["locator"]["page"] == 2 and "Page 2" in b["text"] for b in result.blocks)
    assert result.coverage["units_unparsed"] == 1
    assert any("OCR_REQUIRED" in w and ":2:" in w for w in result.warnings)
    assert (
        hashlib.sha256((CORPUS / document["source_path"]).read_bytes()).hexdigest()
        == document["source_sha256"]
    )


@pytest.mark.parametrize("overlap", [0, 7])
def test_review_long_csv_continuation_keeps_row_key_and_column_alignment(overlap):
    from shopsteward_knowledge.parsing.chunking import chunk_blocks

    result = parse(("SKU,condition\nMilk," + "a" * 1300 + "EXCEPTION\n").encode(), "csv")
    block = result.blocks[0]
    chunks = chunk_blocks(result.blocks, {"max_chars": 1200, "overlap_chars": overlap})
    assert len(chunks) == 2
    assert all(c["text"].startswith("SKU\tcondition\nMilk\t") for c in chunks)
    assert "EXCEPTION" in chunks[1]["text"]
    assert all(len(c["text"]) <= 1200 for c in chunks)
    assert all(c["locator"] == block["locator"] for c in chunks)
    for chunk in chunks:
        start, end = chunk["char_range"]
        assert block["text"][start:end] in chunk["text"]
        assert chunk["content_sha256"] == hashlib.sha256(chunk["text"].encode()).hexdigest()
        assert chunk["table_fragment"]["column"] == 2
        assert chunk["truncated"] is True


def test_review_long_csv_rejects_budget_that_cannot_preserve_row_identity():
    from shopsteward_knowledge.parsing.chunking import chunk_blocks

    result = parse(("SKU,condition\n" + "k" * 1300 + ",EXCEPTION\n").encode(), "csv")
    with pytest.raises(ValueError, match="row context"):
        chunk_blocks(result.blocks, {"max_chars": 1200})


def test_parser_never_opens_a_storage_key_or_accepts_a_path_as_bytes(tmp_path):
    from shopsteward_knowledge.parsing import parse_original

    path = tmp_path / "private.txt"
    path.write_text("private")
    result = parse(b"public", storage_key=str(path))
    assert result.blocks[0]["text"] == "public"
    ref = SimpleNamespace(
        version_id="V", sha256="0" * 64, size_bytes=7, mime="text/plain", storage_key="x"
    )
    with pytest.raises(TypeError):
        parse_original(ref, path)


def test_ooxml_entity_declarations_are_rejected():
    from docx import Document

    stream = io.BytesIO()
    Document().save(stream)
    output = io.BytesIO()
    with ZipFile(stream) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "word/document.xml":
                data = data.replace(
                    b"<w:document",
                    b'<!DOCTYPE x [<!ENTITY ext SYSTEM "file:///no-read">]><w:document',
                )
            target.writestr(item, data)
    with pytest.raises(ValueError, match="declarations"):
        parse(output.getvalue(), "docx")


def test_blank_workbook_chartsheet_remains_explicitly_unparsed():
    from openpyxl import Workbook

    book = Workbook()
    book.create_chartsheet("EmptyChart")
    stream = io.BytesIO()
    book.save(stream)
    result = parse(stream.getvalue(), "xlsx")
    assert result.status == "PARTIAL"
    assert result.blocks == []
    assert any(
        i["kind"] == "chartsheet" and i["status"] == "UNPARSED" for i in result.coverage["items"]
    )


def test_docx_nonbody_evidence_is_reported_as_unparsed():
    from docx import Document

    document = Document()
    document.add_paragraph("Body")
    document.sections[0].footer.paragraphs[0].text = "Binding exception in footer"
    stream = io.BytesIO()
    document.save(stream)
    result = parse(stream.getvalue(), "docx")
    assert result.status == "PARTIAL"
    assert result.blocks[0]["text"] == "Body"
    assert any("FOOTER" in warning for warning in result.warnings)


def test_xlsx_array_formula_is_textual_and_reproducible():
    from openpyxl import Workbook
    from openpyxl.worksheet.formula import ArrayFormula

    book = Workbook()
    book.active.append(["Total"])
    book.active["A2"] = ArrayFormula(ref="A2:A3", text="=SUM(B2:B3*C2:C3)")
    stream = io.BytesIO()
    book.save(stream)
    data = stream.getvalue()
    first, second = parse(data, "xlsx"), parse(data, "xlsx")
    assert first.model_dump() == second.model_dump()
    assert first.blocks[0]["formulas"][0]["formula"] == "=SUM(B2:B3*C2:C3)"
    assert "=SUM(B2:B3*C2:C3)" in first.blocks[0]["text"]


def test_worksheet_embedded_chart_is_not_claimed_complete():
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference

    book = Workbook()
    sheet = book.active
    sheet.append(["Total"])
    sheet.append([5])
    chart = BarChart()
    chart.add_data(Reference(sheet, min_col=1, min_row=1, max_row=2))
    sheet.add_chart(chart, "C1")
    stream = io.BytesIO()
    book.save(stream)
    result = parse(stream.getvalue(), "xlsx")
    assert result.status == "PARTIAL"
    assert result.blocks
    assert any("DRAWING" in warning for warning in result.warnings)


DOCUMENTS = json.loads((CORPUS / "corpus-manifest.json").read_text(encoding="utf-8"))["documents"]


@pytest.mark.parametrize("document", DOCUMENTS, ids=lambda d: d["version_id"])
def test_frozen_k0_originals_preserve_evidence_and_reproduce_ids(document):
    from shopsteward_knowledge.parsing.chunking import chunk_blocks

    path = CORPUS / document["source_path"]
    before = (path.stat().st_mtime_ns, hashlib.sha256(path.read_bytes()).hexdigest())
    data = path.read_bytes()
    result = parse(
        data,
        document["format"],
        version_id=document["version_id"],
        sha256=document["source_sha256"],
        size_bytes=document["source_bytes"],
    )
    from shopsteward_knowledge.contracts import EvidenceLocator

    for block in result.blocks:
        EvidenceLocator.model_validate(block["locator"])
    if document["version_id"] == "K0-S04-v1":
        assert result.status == "OCR_REQUIRED" and result.blocks == []
    elif document["version_id"] == "K0-P02-v1":
        assert result.status == "PARTIAL"
        assert {b["locator"]["page"] for b in result.blocks} == {1}
    else:
        assert result.status in {"COMPLETE", "PARTIAL"}
        assert result.blocks
    text = "".join("".join(b["text"].split()) for b in result.blocks)
    # Gold is only an assertion oracle here; parser receives source bytes only.
    for section in document["sections"]:
        if not section["source_locator"].get("requires_ocr"):
            assert "".join(section["text"].split()) in text
    again = parse(data, document["format"], version_id=document["version_id"])
    assert result.model_dump() == again.model_dump()
    assert chunk_blocks(result.blocks, {}) == chunk_blocks(again.blocks, {})
    assert before == (path.stat().st_mtime_ns, hashlib.sha256(path.read_bytes()).hexdigest())
