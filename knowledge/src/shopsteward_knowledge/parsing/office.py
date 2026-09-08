"""OOXML body/table and worksheet extraction, with explicit missing evidence."""

from __future__ import annotations

from datetime import date, datetime, time
from io import BytesIO
from xml.etree import ElementTree
from zipfile import ZIP_DEFLATED, BadZipFile, ZipFile

from .models import ParseResult, coverage, locator, table_text
from .text import column_name


def _check_archive(content: bytes) -> None:
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 10000 or sum(e.file_size for e in entries) > 100 * 1024 * 1024:
                raise ValueError("Office archive exceeds parser expansion limit")
            for entry in entries:
                if entry.filename.endswith((".xml", ".rels")):
                    xml = archive.read(entry)
                    if b"<!DOCTYPE" in xml.upper() or b"<!ENTITY" in xml.upper():
                        raise ValueError("Office XML declarations/entities are unsupported")
    except BadZipFile as exc:
        raise ValueError("invalid Office archive") from exc


def parse_docx(content: bytes) -> ParseResult:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    _check_archive(content)
    document = Document(BytesIO(content))
    body = document.element.body
    # Count actual XML paragraphs, including empty paragraphs and table cells.
    paragraph_numbers = {element: number for number, element in enumerate(body.iter(qn("w:p")), 1)}
    blocks, items, warnings, sections = [], [], [], []
    table_number = 0
    covered_paragraphs = set()

    def body_elements(container):
        for child in container.iterchildren():
            if child.tag == qn("w:sdt"):
                for content in child.findall(qn("w:sdtContent")):
                    yield from body_elements(content)
            elif child.tag in {qn("w:p"), qn("w:tbl")}:
                yield child

    for element in body_elements(body):
        if element.tag == qn("w:p"):
            covered_paragraphs.add(element)
            paragraph = Paragraph(element, document)
            number = paragraph_numbers[element]
            text = paragraph.text
            style = paragraph.style.name if paragraph.style else ""
            if style.startswith("Heading ") and style[8:].isdigit():
                depth = int(style[8:])
                sections = [(d, t) for d, t in sections if d < depth]
                sections.append((depth, text))
            loc = locator(
                "paragraph", paragraph_range=[number, number], section_path=[t for _, t in sections]
            )
            items.append(dict(locator=loc, status="PARSED"))
            if text.strip():
                blocks.append(dict(text=text, locator=loc))
        elif element.tag == qn("w:tbl"):
            covered_paragraphs.update(element.iter(qn("w:p")))
            table_number += 1
            table = Table(element, document)
            rows = []
            for row in table.rows:
                # Include nested table paragraphs in their containing cell.
                rows.append(
                    [
                        "\n".join(Paragraph(p, cell).text for p in cell._tc.iter(qn("w:p")))
                        for cell in row.cells
                    ]
                )
            headers = rows[0] if rows else []
            data_rows = rows[1:] or ([[]] if rows else [])
            for index, row in enumerate(data_rows, 1):
                source_row = table.rows[index] if len(rows) > 1 else table.rows[0]
                numbers = [paragraph_numbers[p] for p in source_row._tr.iter(qn("w:p"))]
                loc = locator(
                    "table",
                    table=table_number,
                    paragraph_range=[min(numbers), max(numbers)] if numbers else None,
                    section_path=[t for _, t in sections],
                )
                text = table_text(headers, [row] if row else [])
                if text.strip():
                    blocks.append(
                        dict(
                            text=text,
                            locator=loc,
                            table=dict(headers=headers, rows=[row] if row else []),
                        )
                    )
                items.append(dict(locator=loc, status="PARSED"))
    # Other unsupported body wrappers must never silently lose binding text.
    for element, number in paragraph_numbers.items():
        if element not in covered_paragraphs and any(
            node.text and node.text.strip() for node in element.iter(qn("w:t"))
        ):
            warnings.append(f"DOCX_PARAGRAPH_UNPARSED:{number}")
            items.append(
                dict(
                    locator=locator("paragraph", paragraph_range=[number, number]),
                    status="UNPARSED",
                    reason="UNSUPPORTED_CONTAINER",
                )
            )
    # Text extraction is not a claim that embedded visual content was read.
    if any(True for _ in body.iter(qn("w:drawing"))) or any(True for _ in body.iter(qn("w:pict"))):
        warnings.append("DOCX_VISUAL_CONTENT_UNPARSED")
        items.append(dict(kind="visual", status="UNPARSED"))
    with ZipFile(BytesIO(content)) as archive:
        for name in archive.namelist():
            if name.startswith(
                ("word/header", "word/footer", "word/footnotes", "word/endnotes")
            ) and name.endswith(".xml"):
                root = ElementTree.fromstring(archive.read(name))
                if any(node.text and node.text.strip() for node in root.iter(qn("w:t"))):
                    kind = "footer" if name.startswith("word/footer") else "nonbody"
                    warnings.append(f"DOCX_{kind.upper()}_UNPARSED:{name}")
                    items.append(dict(kind=kind, part=name, status="UNPARSED"))
    return ParseResult(
        status="PARTIAL" if warnings else "COMPLETE",
        blocks=blocks,
        warnings=warnings,
        coverage=coverage("body_elements", items),
    )


def _json_value(value):
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    # openpyxl ArrayFormula.__str__ is an object address, not source evidence.
    if isinstance(getattr(value, "text", None), str):
        return value.text
    raise ValueError(f"unsupported worksheet value type: {type(value).__name__}")


def _display(value) -> str:
    value = _json_value(value)
    return "" if value is None else str(value)


def parse_xlsx(content: bytes) -> ParseResult:
    from openpyxl import load_workbook

    _check_archive(content)
    with ZipFile(BytesIO(content)) as archive:
        drawings = [
            name
            for name in archive.namelist()
            if name.startswith("xl/drawings/drawing") and name.endswith(".xml")
        ]
        workbook_xml = ElementTree.fromstring(archive.read("xl/workbook.xml"))
        relationships = ElementTree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        chart_ids = {
            node.get("Id") for node in relationships if node.get("Type", "").endswith("/chartsheet")
        }
        sheets_node = workbook_xml.find("{*}sheets")
        sheet_infos = []
        for node in list(sheets_node):
            is_chart = (
                node.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
                in chart_ids
            )
            sheet_infos.append(
                dict(
                    sheet=node.attrib["name"],
                    state=node.get("state", "visible"),
                    kind="chartsheet" if is_chart else "worksheet",
                )
            )
            if is_chart:
                sheets_node.remove(node)
        # Never ask openpyxl to interpret chart objects. In 3.1 it loses hidden
        # state and raises on empty chartsheets. Retain original chart metadata
        # above, then omit chartsheets only from the in-memory reading copy.
        if chart_ids:
            stream = BytesIO()
            with ZipFile(stream, "w", ZIP_DEFLATED) as target:
                for entry in archive.infolist():
                    data = (
                        ElementTree.tostring(workbook_xml)
                        if entry.filename == "xl/workbook.xml"
                        else archive.read(entry)
                    )
                    target.writestr(entry, data)
            content = stream.getvalue()
    formulas_book = load_workbook(
        BytesIO(content), data_only=False, read_only=True, keep_links=False
    )
    values_book = None
    blocks = []
    items = [dict(kind="drawing", part=name, status="UNPARSED") for name in drawings]
    warnings = [f"XLSX_DRAWING_UNPARSED:{name}" for name in drawings]
    try:
        values_book = load_workbook(
            BytesIO(content), data_only=True, read_only=True, keep_links=False
        )
        for info in sheet_infos:
            name, state = info["sheet"], info["state"]
            if info["kind"] == "chartsheet":
                items.append(info | dict(status="UNPARSED", reason="CHARTSHEET_UNPARSED"))
                warnings.append(f"XLSX_CHARTSHEET_UNPARSED:{name}")
                continue
            sheet = formulas_book[name]
            if (sheet.max_row or 0) * (sheet.max_column or 0) > 1_000_000:
                raise ValueError("worksheet exceeds one million cell iteration limit")
            cache = values_book[name]
            header, header_formulas = None, []
            missing_cache = False
            for row, cached_row in zip(sheet.iter_rows(), cache.iter_rows(), strict=True):
                values = [_display(cell.value) for cell in row]
                if not any(values):
                    continue
                metadata = []
                for cell, cached in zip(row, cached_row, strict=True):
                    if cell.data_type == "f":
                        value = _json_value(cached.value)
                        metadata.append(
                            dict(
                                cell=cell.coordinate,
                                formula=_display(cell.value),
                                cached_value=value,
                                cache_status="missing" if value is None else "present",
                            )
                        )
                        if value is None:
                            missing_cache = True
                            warnings.append(f"XLSX_FORMULA_CACHE_MISSING:{name}!{cell.coordinate}")
                number = next(cell.row for cell in row if cell.value is not None)
                loc = locator(
                    "cells", sheet=name, cell_range=f"A{number}:{column_name(len(row))}{number}"
                )
                if header is None:
                    header, header_formulas = values, metadata
                    header_locator = loc
                    continue
                blocks.append(
                    dict(
                        text=table_text(header, [values]),
                        locator=loc,
                        sheet_state=state,
                        table=dict(headers=header, rows=[values]),
                        formulas=header_formulas + metadata,
                    )
                )
            if header is not None and not any(b["locator"]["sheet"] == name for b in blocks):
                blocks.append(
                    dict(
                        text=table_text(header, []),
                        locator=header_locator,
                        sheet_state=state,
                        table=dict(headers=header, rows=[]),
                        formulas=header_formulas,
                    )
                )
            items.append(info | dict(status="PARTIAL" if missing_cache else "PARSED"))
    finally:
        formulas_book.close()
        if values_book is not None:
            values_book.close()
    return ParseResult(
        status="PARTIAL" if warnings else "COMPLETE",
        blocks=blocks,
        warnings=warnings,
        coverage=coverage("workbook_parts", items),
    )
