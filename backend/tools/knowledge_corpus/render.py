"""Deterministic six-format fixtures and locators checked by actual document extraction.

Office/PDF use the installed bundled document runtime through a persistent local
worker. No installs, network access, or repository dependency changes are needed.
"""

import atexit
import base64
import csv
import io
import json
import os
import re
import subprocess
import sys
import zipfile
from datetime import datetime
from functools import lru_cache
from pathlib import Path

STAMP = datetime(2026, 9, 7)
MARKER = "模拟 SYNTHETIC 商业流程夹具；非真实门店政策；不提供食品安全或法律数值。"
MIMES = {
    "md": "text/markdown",
    "txt": "text/plain",
    "csv": "text/csv",
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}
_worker = None


def _request(payload):
    global _worker
    if _worker is None:
        runtime = os.environ.get("KNOWLEDGE_DOCUMENT_PYTHON") or str(
            Path.home()
            / ".cache/codex-runtimes/codex-primary-runtime/dependencies/python/python.exe"
        )
        if not Path(runtime).is_file():
            raise ValueError("bundled document runtime missing; set KNOWLEDGE_DOCUMENT_PYTHON")
        _worker = subprocess.Popen(
            [runtime, str(Path(__file__).resolve()), "--worker"],
            stdin=subprocess.PIPE,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
            encoding="utf-8",
            creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
        )
        atexit.register(_worker.terminate)
    _worker.stdin.write(json.dumps(payload, ensure_ascii=True) + "\n")
    _worker.stdin.flush()
    line = _worker.stdout.readline()
    if not line:
        raise ValueError("document runtime worker terminated")
    result = json.loads(line)
    if "error" in result:
        raise ValueError("document runtime: " + result["error"])
    return result


def _zip_stable(data):
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(output, "w") as target:
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(name, (2026, 9, 7, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            value = source.read(name)
            if name == "docProps/core.xml":
                value = re.sub(
                    rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)",
                    rb"\g<1>2026-09-07T00:00:00Z\g<2>",
                    value,
                )
            target.writestr(info, value)
    return output.getvalue()


def render_document(title, sections, extension):
    if extension in {"docx", "xlsx", "pdf"}:
        result = _request({"op": "render", "title": title, "sections": sections, "ext": extension})
        return base64.b64decode(result["data"]), result["locators"]
    if extension == "csv":
        stream = io.StringIO(newline="")
        writer = csv.writer(stream, lineterminator="\n")
        writer.writerow(["section", "heading", "text"])
        writer.writerow(["metadata", title, MARKER])
        for i, section in enumerate(sections, 1):
            writer.writerow([i, section["heading"], section["text"]])
        return stream.getvalue().encode("utf-8-sig"), [
            f"row:{i + 3}:C" for i in range(len(sections))
        ]
    text = f"{title}\n{MARKER}\n\n"
    for i, section in enumerate(sections, 1):
        text += f"## section:{i} {section['heading']}\n{section['text']}\n\n"
    return text.encode("utf-8"), [f"section:{i + 1}" for i in range(len(sections))]


@lru_cache(maxsize=2048)
def extract_sections(data, extension):
    if extension in {"docx", "xlsx", "pdf"}:
        return _request(
            {"op": "extract", "data": base64.b64encode(data).decode(), "ext": extension}
        )
    if extension == "csv":
        rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
        return {f"row:{i + 3}:C": row[2] for i, row in enumerate(rows[2:])}
    result = {}
    for block in data.decode("utf-8").split("\n## ")[1:]:
        heading, body = block.split("\n", 1)
        result[heading.split()[0]] = body.strip()
    return result


def _render_binary(title, sections, ext):
    output = io.BytesIO()
    if ext == "docx":
        from docx import Document
        from docx.shared import Pt

        document = Document()
        document.core_properties.created = STAMP
        document.core_properties.modified = STAMP
        document.core_properties.author = "ShopSteward synthetic corpus"
        document.styles["Normal"].font.size = Pt(10)
        document.add_paragraph(title, "Title")
        document.add_paragraph(MARKER)
        for section in sections:
            document.add_heading(section["heading"], 1)
            document.add_paragraph(section["text"])
        document.save(output)
        return _zip_stable(output.getvalue()), [
            f"paragraph:{4 + 2 * i}" for i in range(len(sections))
        ]
    if ext == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font

        workbook = Workbook()
        workbook.properties.created = STAMP
        workbook.properties.modified = STAMP
        sheet = workbook.active
        sheet.title = "条款"
        sheet.append([title, MARKER])
        sheet.append(["section", "heading", "text"])
        for i, section in enumerate(sections, 1):
            sheet.append([i, section["heading"], section["text"]])
            sheet.row_dimensions[i + 2].height = 100
        sheet.column_dimensions["A"].width = 24
        sheet.column_dimensions["B"].width = 40
        sheet.column_dimensions["C"].width = 90
        for row in sheet:
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                cell.font = Font(name="Microsoft YaHei", size=11)
        sheet.freeze_panes = "C3"
        workbook.save(output)
        return _zip_stable(output.getvalue()), [
            f"cells:条款!C{i + 3}" for i in range(len(sections))
        ]
    from xml.sax.saxutils import escape

    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfgen.canvas import Canvas
    from reportlab.platypus import Paragraph

    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    canvas = Canvas(output, pagesize=(595, 842), invariant=1, pageCompression=1)
    canvas.setTitle(title)
    canvas.setAuthor("ShopSteward synthetic corpus")
    style = ParagraphStyle("body", fontName="STSong-Light", fontSize=12, leading=21, wordWrap="CJK")
    for section in sections:
        top = 785
        for text in [title, MARKER, section["heading"], section["text"]]:
            paragraph = Paragraph(escape(text), style)
            _, height = paragraph.wrap(495, 680)
            top -= height
            if top < 70:
                raise ValueError("PDF section exceeds one-page locator budget")
            paragraph.drawOn(canvas, 50, top)
            top -= 22
        canvas.showPage()
    canvas.save()
    return output.getvalue(), [f"page:{i + 1}" for i in range(len(sections))]


def _extract_binary(data, ext):
    if ext == "docx":
        from docx import Document

        return {
            f"paragraph:{i + 1}": p.text
            for i, p in enumerate(Document(io.BytesIO(data)).paragraphs)
        }
    if ext == "xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(data), data_only=True)
        return {
            f"cells:{sheet.title}!{cell.coordinate}": str(cell.value or "")
            for sheet in workbook
            for row in sheet
            for cell in row
        }
    from pypdf import PdfReader

    return {
        f"page:{i + 1}": p.extract_text() for i, p in enumerate(PdfReader(io.BytesIO(data)).pages)
    }


if __name__ == "__main__" and "--worker" in sys.argv:
    for line in sys.stdin:
        try:
            request = json.loads(line)
            if request["op"] == "render":
                data, locators = _render_binary(
                    request["title"], request["sections"], request["ext"]
                )
                result = {"data": base64.b64encode(data).decode(), "locators": locators}
            else:
                result = _extract_binary(base64.b64decode(request["data"]), request["ext"])
        except Exception as error:
            result = {"error": type(error).__name__ + ": " + str(error)}
        print(json.dumps(result, ensure_ascii=True), flush=True)
