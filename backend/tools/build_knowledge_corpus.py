"""Offline K0 fixture builder/validator. No application imports, retrieval or network.

Run from any directory with Python containing reportlab, Pillow, python-docx,
openpyxl and pypdf. Inputs and manually specified gold labels live under
docs/evaluation/knowledge. Only a fixed, owned set of generated paths is written.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import importlib.metadata
import io
import json
import os
import re
import sys
import tempfile
import xml.etree.ElementTree as ET
import zipfile
from collections import Counter
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2] / "docs/evaluation/knowledge"
MARKER = "模拟 / SYNTHETIC — fictional K0 evaluation fixture"
OWNER = "shopsteward-k0-corpus-builder-v1"
STAMP = datetime(2026, 9, 7, 0, 0, 0)
NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


def json_bytes(value):
    return (json.dumps(value, ensure_ascii=False, indent=2) + "\n").encode("utf-8")


def sha(data):
    return hashlib.sha256(data).hexdigest()


def read_json(name):
    return json.loads((ROOT / name).read_text(encoding="utf-8"))


def version_id(doc):
    return f"K0-{doc['key']}-v{doc.get('version', 1)}"


def store_id(doc):
    return "E-STORE" if doc["store"] == "east" else "W-STORE"


def document_id(doc):
    return f"K0-{'E' if doc['store'] == 'east' else 'W'}-{doc.get('document', doc['key'])}"


def canonical_sections(doc):
    return [
        {"section_id": f"section:{i}", "heading": h, "text": t}
        for i, (h, t) in enumerate(doc["sections"], 1)
    ]


def normalized_zip(data, patch=None):
    """Pin ZIP timestamps/order; optionally supply exact XLSX formula caches."""
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source:
        with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
            for name in sorted(source.namelist()):
                value = source.read(name)
                if patch:
                    value = patch(name, value)
                info = zipfile.ZipInfo(name, (2026, 9, 7, 0, 0, 0))
                info.compress_type = zipfile.ZIP_DEFLATED
                info.external_attr = 0o600 << 16
                target.writestr(info, value)
    return output.getvalue()


def source_rows(doc):
    return [(s["section_id"], s["heading"], s["text"]) for s in canonical_sections(doc)]


def make_docx(doc):
    from docx import Document

    document = Document()
    document.core_properties.created = STAMP
    document.core_properties.modified = STAMP
    document.core_properties.author = "ShopSteward fictional K0 fixtures"
    document.add_paragraph(MARKER)
    document.add_heading(doc["title"], 0)
    document.add_paragraph(f"{store_id(doc)} | {version_id(doc)}")
    for i, (_, heading, text) in enumerate(source_rows(doc), 1):
        document.add_heading(f"[section:{i}] {heading}", 1)
        document.add_paragraph(text)
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "定位 / Locator"
    table.rows[0].cells[1].text = "操作核对 / Operational check"
    # Structured recap remains part of the same authored document, not extra corpus.
    for section, heading, text in source_rows(doc):
        row = table.add_row().cells
        row[0].text = f"{section} {heading}"
        row[1].text = text
    stream = io.BytesIO()
    document.save(stream)
    return normalized_zip(stream.getvalue())


FORMULAS = {
    "S07": (12, 50, "=B3*C3", 600, "袋/箱 × 只/袋 = 只/箱"),
    "P04": (8, 125, "=B3*C3", 1000, "杯/箱 × g/杯 = g/箱"),
    "A03": (6, 20, "=B3*C3", 120, "桌 × 只/桌 = 只"),
    "R03": (2, 10, "=B3/C3", 0.2, "漏贴项 / 巡检项"),
}


def make_xlsx(doc):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment

    wb = Workbook()
    wb.properties.created = STAMP
    wb.properties.modified = STAMP
    ws = wb.active
    ws.title = "说明"
    ws.append([MARKER])
    ws.append(["section_id", "heading", "text"])
    for row in source_rows(doc):
        ws.append(row)
    for col, width in [("A", 20), ("B", 34), ("C", 100)]:
        ws.column_dimensions[col].width = width
    for row in range(3, 6):
        ws.row_dimensions[row].height = 70
        ws.cell(row, 3).alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "C3"
    calc = wb.create_sheet("换算")
    calc.append([MARKER])
    calc.append(["单位/units", "输入/input A", "输入/input B", "公式/formula"])
    a, b, formula, result, units = FORMULAS[doc["key"]]
    calc.append([units, a, b, formula])
    calc.column_dimensions["A"].width = 40
    for col in "BCD":
        calc.column_dimensions[col].width = 22
    wb.calculation.fullCalcOnLoad = True
    stream = io.BytesIO()
    wb.save(stream)

    def patch(name, value):
        if name == "docProps/core.xml":
            value = re.sub(
                rb"(<dcterms:modified[^>]*>).*?(</dcterms:modified>)",
                rb"\g<1>2026-09-07T00:00:00Z\2",
                value,
            )
        if name == "xl/worksheets/sheet2.xml" and doc["formula_cache"]:
            xml = ET.fromstring(value)
            cell = xml.find(".//s:c[@r='D3']", NS)
            cache = cell.find("s:v", NS)
            if cache is None:
                cache = ET.SubElement(cell, "{" + NS["s"] + "}v")
            cache.text = str(result)
            value = ET.tostring(xml, encoding="utf-8", xml_declaration=True)
        return value

    return normalized_zip(stream.getvalue(), patch)


def wrap_chars(text, font, limit):
    lines, line = [], ""
    for char in text:
        if font.getlength(line + char) > limit and line:
            lines.append(line)
            line = ""
        line += char
    if line:
        lines.append(line)
    return lines


def make_pdf(doc, font_path):
    from PIL import Image, ImageDraw, ImageFont
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfgen import canvas

    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    stream = io.BytesIO()
    c = canvas.Canvas(stream, pagesize=(595, 842), invariant=1, pageCompression=1)
    c.setTitle(f"SYNTHETIC {version_id(doc)}")
    c.setAuthor("ShopSteward K0 fictional fixtures")
    mode = doc.get("pdf_mode", "text")

    def text_page(rows, page):
        c.setFont("STSong-Light", 11)
        y = 805
        for line in [MARKER, doc["title"], f"{version_id(doc)} | {store_id(doc)} | page {page}"]:
            c.drawString(32, y, line)
            y -= 23
        for heading, content in rows:
            for line in [heading, *[content[i : i + 48] for i in range(0, len(content), 48)]]:
                c.drawString(32, y, line)
                y -= 20
            y -= 18

    def scan_page(rows, page):
        # Image-only page: no invisible text/OCR overlay and no answer side channel.
        bitmap = Image.new("RGB", (1190, 1684), "#f6f4ef")
        draw = ImageDraw.Draw(bitmap)
        font = ImageFont.truetype(str(font_path), 23)
        y = 65
        for text in [MARKER, doc["title"], f"{version_id(doc)} | page {page}"]:
            draw.text((64, y), text, font=font, fill="black")
            y += 44
        for heading, content in rows:
            for line in [heading, *wrap_chars(content, font, 1040)]:
                draw.text((64, y), line, font=font, fill="#242424")
                y += 38
            y += 35
        c.drawImage(ImageReader(bitmap), 0, 0, width=595, height=842)

    rows = [(f"[{s}] {h}", t) for s, h, t in source_rows(doc)]
    if mode == "scan":
        scan_page(rows, 1)
        c.showPage()
    elif mode == "mixed":
        text_page(rows[:1], 1)
        c.showPage()
        scan_page(rows[1:], 2)
        c.showPage()
    elif mode == "multipage_table":
        checks = [
            ("01", "收货员", "供应商匹配", "供应商ID"),
            ("02", "收货员", "订单号匹配", "订单号"),
            ("03", "收货员", "批号可读", "批号照片"),
            ("04", "收货员", "通道畅通", "到店时间"),
            ("05", "收货员", "外箱无渗漏", "外箱照片"),
            ("06", "收货员", "逐箱数量核对", "实收数量"),
            ("07", "收货员", "冷品测温", "温度记录"),
            ("08", "收货员", "异常转隔离", "隔离标签"),
            ("09", "司机", "差异签认", "司机签字"),
            ("10", "收货员", "保留待交余额", "差异单"),
            ("11", "店长", "联系供应商", "补发回执"),
            ("12", "店长", "复核差异单", "复核签字"),
        ]
        for page in range(1, 4):
            text_page(rows[page - 1 : page], page)
            xs = [32, 95, 185, 390, 563]
            y = 540
            table = [("步骤", "责任人", "判定", "记录"), *checks[(page - 1) * 4 : page * 4]]
            for row in table:
                for x in xs:
                    c.line(x, y, x, y - 48)
                c.line(32, y, 563, y)
                for index, value in enumerate(row):
                    c.drawString(xs[index] + 7, y - 28, value)
                y -= 48
            c.line(32, y, 563, y)
            c.drawString(32, 250, f"[section:3/table] continued checklist {page}/3")
            c.showPage()
    else:
        text_page(rows, 1)
        c.showPage()
    c.save()
    return stream.getvalue()


def source_bytes(doc, font_path):
    fmt = doc["format"]
    if fmt == "pdf":
        return make_pdf(doc, font_path)
    if fmt == "docx":
        return make_docx(doc)
    if fmt == "xlsx":
        return make_xlsx(doc)
    if fmt == "csv":
        stream = io.StringIO(newline="")
        writer = csv.writer(stream, lineterminator="\n")
        writer.writerow(
            ["synthetic_notice", "store_id", "version_id", "section_id", "heading", "text"]
        )
        for section, heading, text in source_rows(doc):
            writer.writerow([MARKER, store_id(doc), version_id(doc), section, heading, text])
        return stream.getvalue().encode("utf-8-sig")
    lines = [MARKER, f"# {doc['title']}", f"{store_id(doc)} | {version_id(doc)}"]
    for section, heading, text in source_rows(doc):
        lines += ["", f"## [{section}] {heading}", text]
    return ("\n".join(lines) + "\n").encode("utf-8")


def source_locator(doc, i):
    fmt, mode = doc["format"], doc.get("pdf_mode", "text")
    if fmt == "pdf":
        page = i if mode == "multipage_table" else (2 if mode == "mixed" and i > 1 else 1)
        return {
            "type": "pdf",
            "page": page,
            "anchor": f"[section:{i}]",
            "requires_ocr": mode == "scan" or (mode == "mixed" and page == 2),
        }
    if fmt == "xlsx":
        return {"type": "xlsx", "sheet": "说明", "range": f"A{i + 2}:C{i + 2}"}
    if fmt == "csv":
        return {"type": "csv", "row": i + 1, "columns": ["heading", "text"]}
    if fmt == "docx":
        return {
            "type": "docx",
            "heading": f"[section:{i}] {doc['sections'][i - 1][0]}",
            "paragraph_index": 2 * i + 3,
        }
    return {"type": fmt, "heading": f"## [section:{i}] {doc['sections'][i - 1][0]}"}


def prepare(font_path):
    spec = read_json("corpus-spec.json")
    outputs, entries = {}, []
    for doc in spec["documents"]:
        vid = version_id(doc)
        path = f"sources/{store_id(doc)}/{vid}/{doc['title']}.{doc['format']}"
        raw = source_bytes(doc, font_path)
        sections = canonical_sections(doc)
        lines = [MARKER, doc["title"], f"{store_id(doc)} | {vid}"]
        for index, section in enumerate(sections, 1):
            section["source_locator"] = source_locator(doc, index)
            section["canonical_start_line"] = len(lines) + 1
            lines += [f"[{section['section_id']}] {section['heading']}", section["text"]]
            section["canonical_end_line"] = len(lines)
        canonical = ("\n".join(lines) + "\n").encode("utf-8")
        canonical_path = f"canonical/{vid}.txt"
        warnings = []
        if doc.get("pdf_mode") == "scan":
            warnings = [{"code": "ocr_required", "pages": [1], "without_ocr": "no_text"}]
        elif doc.get("pdf_mode") == "mixed":
            warnings = [
                {"code": "partial_text_ocr_required", "pages": [2], "without_ocr": "partial"}
            ]
        elif doc.get("pdf_mode") == "multipage_table":
            warnings = [
                {
                    "code": "cross_page_table",
                    "pages": [1, 2, 3],
                    "expectation": "repeat headers; flag truncation if table continuation is lost",
                }
            ]
        if doc["format"] == "xlsx" and not doc["formula_cache"]:
            warnings = [
                {
                    "code": "formula_cache_missing",
                    "sheet": "换算",
                    "cells": ["D3"],
                    "without_calculation": "no_cached_value",
                }
            ]
        entry = {
            "key": doc["key"],
            "synthetic": True,
            "store_id": store_id(doc),
            "document_id": document_id(doc),
            "version_id": vid,
            "version_number": doc.get("version", 1),
            "category": doc["category"],
            "title": doc["title"],
            "language": ["zh-CN", "en"],
            "source_path": path,
            "format": doc["format"],
            "source_sha256": sha(raw),
            "source_bytes": len(raw),
            "canonical_text_path": canonical_path,
            "canonical_sha256": sha(canonical),
            "canonical_role": "evaluation_only_transcription_not_parser_output",
            "valid_from": doc.get("valid_from", "2026-09-01"),
            "valid_until": doc.get("valid_until", "2026-12-01"),
            "status": doc.get("status", "published"),
            "allowed_roles": doc.get("allowed_roles", ["staff", "manager"]),
            "expected_parser_warnings": warnings,
            "sections": sections,
            "conflict_set": doc.get("conflict_set"),
        }
        if doc["format"] == "xlsx":
            entry["formula_fixture"] = {
                "sheet": "换算",
                "cell": "D3",
                "formula": FORMULAS[doc["key"]][2],
                "has_cache": doc["formula_cache"],
            }
        entries.append(entry)
        outputs[path] = raw
        outputs[canonical_path] = canonical
    outputs["corpus-manifest.json"] = json_bytes(
        {
            "schema_version": "1.0",
            "generator": OWNER,
            "synthetic": True,
            "as_of": spec["as_of"],
            "interval_semantics": "[valid_from, valid_until)",
            "source_scope": (
                "Only sources/ files are parser inputs. "
                "Never index labels, canonical transcriptions or evaluation reports."
            ),
            "stores": [
                {"store_id": "E-STORE", "name": "模拟东店 / Synthetic East"},
                {"store_id": "W-STORE", "name": "模拟西店 / Synthetic West"},
            ],
            "documents": entries,
        }
    )
    return outputs


def safe_path(relative):
    path = ROOT / relative
    if not path.resolve().is_relative_to(ROOT.resolve()) or path.is_symlink():
        raise ValueError(f"Unsafe output path: {relative}")
    return path


def write_owned(outputs):
    inventory = ROOT / "generated-files.json"
    prior = (
        read_json("generated-files.json")
        if inventory.exists()
        else {"generator": OWNER, "files": {}}
    )
    if prior["generator"] != OWNER:
        raise ValueError("Refusing to use an inventory owned by another generator")
    # Validate all targets before writing anything; refuse unknown existing files.
    for name in outputs:
        path = safe_path(name)
        if path.exists() and name not in prior["files"]:
            raise FileExistsError(f"Refusing to overwrite unowned file: {name}")
        if path.exists() and sha(path.read_bytes()) != prior["files"].get(name):
            raise ValueError(f"Owned output was edited externally; refusing overwrite: {name}")
    inventory_data = {
        "generator": OWNER,
        "files": {**prior["files"], **{k: sha(v) for k, v in outputs.items()}},
    }
    for name, value in [*outputs.items(), ("generated-files.json", json_bytes(inventory_data))]:
        path = safe_path(name)
        path.parent.mkdir(parents=True, exist_ok=True)
        if path.exists() and path.read_bytes() == value:
            continue
        # Atomic replacement of this one owned file, no directory deletion.
        with tempfile.NamedTemporaryFile(dir=path.parent, prefix=".k0-build-", delete=False) as tmp:
            temp_path = Path(tmp.name)
            tmp.write(value)
        try:
            os.replace(temp_path, path)
        finally:
            if temp_path.exists():
                temp_path.unlink()


def verify():
    from docx import Document
    from openpyxl import load_workbook
    from pypdf import PdfReader

    manifest = read_json("corpus-manifest.json")
    entries = manifest["documents"]
    versions = {d["version_id"]: d for d in entries}
    checks = []

    def check(condition, message):
        if not condition:
            raise ValueError(message)
        checks.append(message)

    check(len(entries) == len(versions) == 40, "40 unique document versions")
    check(
        Counter(d["category"] for d in entries)
        == {"supplier": 12, "product": 8, "sop": 8, "campaign": 6, "retrospective": 6},
        "category counts 12/8/8/6/6",
    )
    counts = Counter(d["format"] for d in entries)
    check(
        all(counts[f] >= n for f, n in [("pdf", 8), ("docx", 6), ("xlsx", 4), ("csv", 4)]),
        "minimum format counts",
    )
    check(len({d["source_sha256"] for d in entries}) == 40, "40 distinct source hashes")
    check(
        len({tuple(s["text"] for s in d["sections"]) for d in entries}) == 40,
        "40 distinct authored section sets",
    )
    file_checks = []
    for doc in entries:
        path = safe_path(doc["source_path"])
        check(sha(path.read_bytes()) == doc["source_sha256"], f"source hash {doc['version_id']}")
        canonical = safe_path(doc["canonical_text_path"]).read_bytes()
        check(sha(canonical) == doc["canonical_sha256"], f"canonical hash {doc['version_id']}")
        check(MARKER in canonical.decode("utf-8"), f"simulation marker {doc['version_id']}")
        observed = {"version_id": doc["version_id"], "format": doc["format"]}
        fmt = doc["format"]
        if fmt == "pdf":
            reader = PdfReader(path)
            text = [p.extract_text() or "" for p in reader.pages]
            images = [len(p.images) for p in reader.pages]
            observed.update(pages=len(text), text_characters=list(map(len, text)), images=images)
            if doc["key"] == "S04":
                check(
                    len(text) == 1 and not text[0].strip() and images[0] > 0,
                    "scan PDF has image and no text layer",
                )
            elif doc["key"] == "P02":
                check(
                    len(text) == 2 and len(text[0]) > 100 and not text[1].strip() and images[1] > 0,
                    "mixed PDF has text page plus scan page",
                )
            elif doc["key"] == "O01":
                check(
                    len(text) == 3 and all("责任人" in p for p in text) and "复核差异单" in text[2],
                    "3-page table has repeated headers and final row",
                )
            else:
                check(MARKER in text[0], f"text PDF marker {doc['version_id']}")
            # Canonical text is ground truth; these checks are artifact integrity, not OCR scoring.
            for section in doc["sections"]:
                loc = section["source_locator"]
                if not loc["requires_ocr"]:
                    check(
                        loc["anchor"] in text[loc["page"] - 1],
                        f"PDF source anchor {doc['version_id']} {section['section_id']}",
                    )
        elif fmt == "docx":
            document = Document(path)
            observed.update(paragraphs=len(document.paragraphs), tables=len(document.tables))
            check(
                len(document.tables) == 1 and document.paragraphs[0].text == MARKER,
                f"DOCX structure {doc['version_id']}",
            )
            for section in doc["sections"]:
                loc = section["source_locator"]
                # paragraph_index is one-based and points to the body after its heading.
                check(
                    document.paragraphs[loc["paragraph_index"] - 1].text == section["text"],
                    f"DOCX locator {doc['version_id']} {section['section_id']}",
                )
        elif fmt == "xlsx":
            formula_wb = load_workbook(path, data_only=False)
            data_wb = load_workbook(path, data_only=True)
            cached = data_wb["换算"]["D3"].value
            observed.update(formula=formula_wb["换算"]["D3"].value, cached_value=cached)
            check(
                observed["formula"] == doc["formula_fixture"]["formula"],
                f"XLSX formula {doc['version_id']}",
            )
            expected = FORMULAS[doc["key"]][3] if doc["formula_fixture"]["has_cache"] else None
            check(cached == expected, f"XLSX cache state {doc['version_id']}")
            check(data_wb["说明"]["A1"].value == MARKER, f"XLSX marker {doc['version_id']}")
            for i, section in enumerate(doc["sections"], 3):
                check(
                    data_wb["说明"].cell(i, 3).value == section["text"],
                    f"XLSX locator {doc['version_id']} {section['section_id']}",
                )
            formula_wb.close()
            data_wb.close()
        elif fmt == "csv":
            with path.open(encoding="utf-8-sig", newline="") as stream:
                rows = list(csv.DictReader(stream))
            check(
                len(rows) == 3 and all(r["synthetic_notice"] == MARKER for r in rows),
                f"CSV rows {doc['version_id']}",
            )
            check(
                [r["text"] for r in rows] == [s["text"] for s in doc["sections"]],
                f"CSV locators {doc['version_id']}",
            )
        else:
            check(
                MARKER in path.read_text(encoding="utf-8"),
                f"text source marker {doc['version_id']}",
            )
        file_checks.append(observed)

    fixture = read_json("entities-fixture.json")
    nodes = {n["entity_id"]: n for n in fixture["entities"]}
    records = {r["record_id"]: r for r in fixture["backend_records"]}
    edges = read_json("relations.json")["edges"]
    edge_map = {e["edge_id"]: e for e in edges}
    check(len(edge_map) == len(edges), "unique relation IDs")
    for edge in edges:
        for side in ("subject", "object"):
            eid = edge[f"{side}_id"]
            if edge[f"{side}_type"] == "DocumentVersion":
                check(
                    eid in versions and versions[eid]["store_id"] == edge["store_id"],
                    f"relation document scope {edge['edge_id']} {side}",
                )
            else:
                check(
                    eid in nodes
                    and nodes[eid]["store_id"] == edge["store_id"]
                    and nodes[eid]["entity_type"] == edge[f"{side}_type"],
                    f"relation entity scope {edge['edge_id']} {side}",
                )
        check(
            edge["valid_from"] < edge["valid_until"] and edge["revision"] >= 1,
            f"relation validity {edge['edge_id']}",
        )
        check(
            edge["assertion_status"] in ["verified", "candidate", "rejected"]
            and isinstance(edge["qualifiers"], dict),
            f"relation status and conditions {edge['edge_id']}",
        )
        if edge["origin_kind"] == "backend":
            record = records[edge["source_record_id"]]
            check(
                record["store_id"] == edge["store_id"]
                and all(
                    record[k] == edge[k]
                    for k in [
                        "subject_type",
                        "subject_id",
                        "predicate",
                        "object_type",
                        "object_id",
                        "valid_from",
                        "valid_until",
                    ]
                ),
                f"backend relation provenance {edge['edge_id']}",
            )
        else:
            src = versions[edge["source_version_id"]]
            check(
                src["store_id"] == edge["store_id"]
                and edge["source_locator"] in {s["section_id"] for s in src["sections"]},
                f"document relation provenance {edge['edge_id']}",
            )

    summaries = {}
    for filename, total, splits, groups in [
        (
            "cases.jsonl",
            60,
            {"dev": 40, "test": 20},
            {
                "exact": 12,
                "keyword": 12,
                "synonym": 12,
                "synthesis": 8,
                "no_answer": 8,
                "business_mixed": 8,
            },
        ),
        (
            "relation-cases.jsonl",
            12,
            {"dev": 8, "test": 4},
            {"single_step": 4, "multi_step": 4, "conditional": 2, "unknown": 2},
        ),
    ]:
        cases = [
            json.loads(line)
            for line in (ROOT / filename).read_text(encoding="utf-8").splitlines()
            if line.strip()
        ]
        check(
            len(cases) == total and len({c["id"] for c in cases}) == total,
            f"{filename} count and unique IDs",
        )
        check(Counter(c["split"] for c in cases) == splits, f"{filename} split counts")
        check(Counter(c["group"] for c in cases) == groups, f"{filename} group counts")
        check(len({c["query"] for c in cases}) == total, f"{filename} unique queries")
        for case in cases:
            check(
                bool(case["required_assertions"]) and bool(case["forbidden_assertions"]),
                f"label assertions {case['id']}",
            )
            principal = fixture["principals"][case["store_fixture"]]
            for vid in case["expected_document_versions"]:
                d = versions[vid]
                check(
                    d["store_id"] == principal["store_id"]
                    and principal["role"] in d["allowed_roles"],
                    f"gold ACL {case['id']} {vid}",
                )
                if not case.get("allow_historical", False):
                    check(
                        d["status"] == "published"
                        and d["valid_from"] <= case["as_of"] < d["valid_until"],
                        f"gold time {case['id']} {vid}",
                    )
            for evidence in case.get("required_evidence", []):
                check(
                    evidence["version_id"] in case["expected_document_versions"]
                    and set(evidence["sections"])
                    <= {s["section_id"] for s in versions[evidence["version_id"]]["sections"]},
                    f"gold locator {case['id']}",
                )
            for eid in (
                case.get("expected_edge_ids", [])
                + case.get("excluded_edge_ids", [])
                + case.get("diagnostic_edge_ids", [])
            ):
                check(eid in edge_map, f"gold edge exists {case['id']} {eid}")
            for path in case.get("expected_paths", []):
                path_edges = [edge_map[hop["edge_id"]] for hop in path]
                endpoints = [
                    (e["subject_id"], e["object_id"])
                    if hop["direction"] == "forward"
                    else (e["object_id"], e["subject_id"])
                    for e, hop in zip(path_edges, path, strict=False)
                ]
                check(
                    all(a[1] == b[0] for a, b in zip(endpoints, endpoints[1:], strict=False)),
                    f"gold path contiguous {case['id']}",
                )
                check(
                    all(
                        e["store_id"] == principal["store_id"]
                        and e["assertion_status"] == "verified"
                        and e["valid_from"] <= case["as_of"] < e["valid_until"]
                        for e in path_edges
                    ),
                    f"gold path scope/status/time {case['id']}",
                )
                plan = case["query_plan"]
                check(
                    endpoints[0][0] in plan["seed_entity_ids"] and len(path) <= plan["max_depth"],
                    f"gold path matches query seed/depth {case['id']}",
                )
                check(
                    all(e["predicate"] in plan["allowed_predicates"] for e in path_edges),
                    f"gold path allowed predicates {case['id']}",
                )
                final_id = endpoints[-1][1]
                final_type = (
                    "DocumentVersion" if final_id in versions else nodes[final_id]["entity_type"]
                )
                check(
                    final_type in plan["target_types"]
                    and (not plan["target_entity_ids"] or final_id in plan["target_entity_ids"]),
                    f"gold path target {case['id']}",
                )
            if filename == "relation-cases.jsonl":
                check(
                    case["facts_input"] == "relations.json" and case["comparators"] == ["C", "D"],
                    f"same C/D facts {case['id']}",
                )
                plan = case["query_plan"]
                check(
                    bool(plan["seed_entity_ids"])
                    and all(s in nodes or s in versions for s in plan["seed_entity_ids"]),
                    f"independent query seed exists {case['id']}",
                )
                check(
                    1 <= plan["max_depth"] <= 3
                    and plan["direction"] in ["forward", "reverse", "both"]
                    and isinstance(plan["qualifier_context"], dict),
                    f"bounded independent query intent {case['id']}",
                )
        summaries[filename] = {
            "count": len(cases),
            "splits": dict(Counter(c["split"] for c in cases)),
            "groups": dict(Counter(c["group"] for c in cases)),
            "sha256": sha((ROOT / filename).read_bytes()),
        }
    for name, digest in read_json("generated-files.json")["files"].items():
        check(sha(safe_path(name).read_bytes()) == digest, f"ownership hash {name}")
    return {
        "schema_version": "1.0",
        "as_of": "2026-09-07",
        "status": "passed",
        "scope": (
            "Offline artifact and label-reference integrity only; "
            "no retrieval, OCR, embedding, database or external API evaluation."
        ),
        "human_review_status": (
            "pending; gold labels are individually specified by the authoring assistant, "
            "not retriever-generated and not claimed human-reviewed"
        ),
        "source_count": len(entries),
        "logical_document_count": len({d["document_id"] for d in entries}),
        "formats": dict(counts),
        "categories": dict(Counter(d["category"] for d in entries)),
        "cases": summaries,
        "relation_count": len(edges),
        "entity_count": len(nodes),
        "checks_passed": len(checks),
        "files": file_checks,
        "dependencies": {
            name: importlib.metadata.version(name)
            for name in ["reportlab", "pillow", "python-docx", "openpyxl", "pypdf"]
        },
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--verify-only",
        action="store_true",
        help="Read-only integrity validation; never writes files",
    )
    parser.add_argument(
        "--check-reproducible",
        action="store_true",
        help="Rebuild sources in memory and compare every generated byte",
    )
    parser.add_argument(
        "--font",
        type=Path,
        default=Path(os.environ.get("K0_CJK_FONT", "C:/Windows/Fonts/msyh.ttc")),
        help="CJK TrueType/OpenType font for image-only PDF pages",
    )
    args = parser.parse_args()
    if not args.verify_only:
        if not args.font.is_file():
            parser.error(
                "A local CJK font is required; pass --font or K0_CJK_FONT. "
                "No font download is performed."
            )
        outputs = prepare(args.font)
        if args.check_reproducible:
            for name, value in outputs.items():
                if safe_path(name).read_bytes() != value:
                    raise ValueError(f"Rebuild differs: {name}")
        else:
            write_owned(outputs)
    report = verify()
    if not args.verify_only and not args.check_reproducible:
        # The report excludes its own inventory entry from the check count, so repeated
        # builds produce the same report. All inventory hashes are still checked.
        if "validation-report.json" in read_json("generated-files.json")["files"]:
            report["checks_passed"] -= 1
        write_owned({"validation-report.json": json_bytes(report)})
    print(
        json.dumps(
            {
                "status": report["status"],
                "sources": report["source_count"],
                "formats": report["formats"],
                "cases": {k: v["splits"] for k, v in report["cases"].items()},
                "relations": report["relation_count"],
                "checks_passed": report["checks_passed"],
                "reproducible": bool(args.check_reproducible),
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    try:
        main()
    except (ValueError, FileExistsError) as exc:
        print(f"K0 validation failed: {exc}", file=sys.stderr)
        sys.exit(1)
